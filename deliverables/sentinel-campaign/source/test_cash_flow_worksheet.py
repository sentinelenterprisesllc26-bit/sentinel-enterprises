"""Tests for the Sentinel cash-flow worksheet generator."""

from pathlib import Path
import tempfile

import openpyxl

from generate_cash_flow_worksheet import (
    DAILY_ROWS,
    DEFAULT_OUTPUT,
    cash_net_received,
    generate_workbook,
    operating_net_sales_estimate,
    safe_ratio,
    scenario_gross,
    scenario_orders_required,
)


def test_generator_creates_expected_sheets_and_blank_actual_inputs(tmp_path: Path) -> None:
    output = generate_workbook(tmp_path / "cash-flow-worksheet.xlsx")
    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook.sheetnames == [
        "Start Here",
        "Daily Inputs",
        "Activity Ledger",
        "Cash Ledger",
        "Scenario",
        "Weekly Review",
        "Glossary",
    ]
    inputs = workbook["Daily Inputs"]
    # All actual input cells are intentionally blank, including dates.
    for row in range(6, 6 + DAILY_ROWS):
        assert all(inputs.cell(row=row, column=col).value is None for col in range(1, 16))
    assert DEFAULT_OUTPUT.name == "cash-flow-worksheet.xlsx"


def test_formula_boundaries_preserve_unknowns_and_cash_boundary(tmp_path: Path) -> None:
    output = generate_workbook(tmp_path / "cash-flow-worksheet.xlsx")
    workbook = openpyxl.load_workbook(output, data_only=False)
    activity = workbook["Activity Ledger"]
    cash = workbook["Cash Ledger"]
    inputs = workbook["Daily Inputs"]
    assert activity["F6"].value == '=IF(OR(E6="",D6="",D6<=0),"",E6/D6)'
    assert activity["G6"].value == '=IF(OR(D6="",C6="",C6<=0),"",D6/C6)'
    assert activity["J6"].value == '=IF(OR(E6="",H6="",I6=""),"",E6+H6-I6)'
    assert activity["N6"].value == '=IF(OR(J6="",K6="",L6="",M6=""),"",J6-K6-L6-M6)'
    assert cash["G6"].value == '=IF(OR(B6="",C6="",D6="",E6="",F6=""),"",B6+C6-D6-E6-F6)'
    assert inputs["M5"].value == "Other cash expenses paid"
    assert inputs["O5"].value == "Refunds paid separately (NOT deducted from payout)"
    assert cash["D5"].value == "Refunds paid separately (NOT deducted from payout)"
    # There is no processor-fee subtraction or activity-expense subtraction in the cash formula.
    assert "Processor fees" not in cash["G6"].value
    assert workbook.calculation.fullCalcOnLoad is True


def test_all_zero_aggregate_rates_are_blank_safe(tmp_path: Path) -> None:
    output = generate_workbook(tmp_path / "cash-flow-worksheet.xlsx")
    workbook = openpyxl.load_workbook(output, data_only=False)
    activity = workbook["Activity Ledger"]
    # All-zero known orders/clicks must not produce #DIV/0! in aggregate AOV/conversion.
    assert "SUM(D6:D19)>0" in activity["F20"].value
    assert "SUM(C6:C19)>0" in activity["G20"].value
    assert safe_ratio(0, 0) is None


def test_netted_refund_and_separate_refund_cash_boundaries() -> None:
    # A payout that already nets the refund uses zero in the separate-refund field.
    assert cash_net_received(100, 5, 0, 3, 2) == 100
    # A refund paid outside the payout is deducted once, and only once.
    assert cash_net_received(100, 5, 10, 3, 2) == 90
    assert cash_net_received(100, 5, None, 3, 2) is None


def test_scenario_sample_and_blank_optional_conversion() -> None:
    assert scenario_orders_required(100.00, 17.99) == 6
    assert round(scenario_gross(100.00, 17.99), 2) == 107.94
    assert safe_ratio(None, 10) is None
    assert safe_ratio(3, 0) is None
    assert safe_ratio(3, 12) == 0.25
    assert operating_net_sales_estimate(100, None, 0) is None
    assert operating_net_sales_estimate(100, 10, 5) == 105


def test_scenario_has_explicit_illustrative_values_and_blank_conversion(tmp_path: Path) -> None:
    output = generate_workbook(tmp_path / "cash-flow-worksheet.xlsx")
    workbook = openpyxl.load_workbook(output, data_only=True)
    scenario = workbook["Scenario"]
    assert scenario["B5"].value == 100
    assert scenario["B6"].value == 17.99
    assert scenario["B7"].value == 6
    assert scenario["B8"].value == 107.94
    # Optional conversion is unknown and therefore blank, not a fabricated result.
    assert scenario["B10"].value is None
    assert scenario["B11"].value is None


def test_scenario_funnel_stage_matches_named_rate(tmp_path: Path) -> None:
    output = generate_workbook(tmp_path / "cash-flow-worksheet.xlsx")
    workbook = openpyxl.load_workbook(output, data_only=False)
    scenario = workbook["Scenario"]
    rate_label = scenario["A10"].value.lower()
    output_label = scenario["A11"].value.lower()
    assert "checkout" in rate_label and "paid-order" in rate_label and "rate" in rate_label
    assert "checkout clicks required" in output_label
    assert "visitors needed" not in output_label
    # Required orders divided by checkout-to-paid-order rate yields checkout clicks.
    assert scenario["B11"].value == '=IF(OR(B10="",B10<=0),"",B7/B10)'


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory)
        test_generator_creates_expected_sheets_and_blank_actual_inputs(path)
        test_formula_boundaries_preserve_unknowns_and_cash_boundary(path)
        test_all_zero_aggregate_rates_are_blank_safe(path)
        test_netted_refund_and_separate_refund_cash_boundaries()
        test_scenario_sample_and_blank_optional_conversion()
        test_scenario_has_explicit_illustrative_values_and_blank_conversion(path)
        test_scenario_funnel_stage_matches_named_rate(path)
    print("All cash-flow worksheet tests passed.")